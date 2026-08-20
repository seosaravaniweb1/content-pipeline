"""ورودی/خروجی گوگل‌شیت محصولات (فاز ۵).

این ماژول تنها جایی است که به شیت شما دست می‌زند، و با سه قاعده‌ی سفت‌وسخت
نوشته شده چون اشتباهش روی ۱۵ هزار ردیف قابل برگشت نیست:

1. **فقط ستون‌های نگاشت‌شده نوشته می‌شوند.** ستون‌های «وضعیت» و
   «شناسه محصول» مالِ اسکریپت درج محصول‌اند و در فهرست ``never_write``
   می‌مانند؛ هیچ‌وقت مقداری در آن‌ها نوشته نمی‌شود.
2. **سلول پرشده دست نمی‌خورد** (مگر ``overwrite: always``). چیزی که خودتان
   دستی نوشته‌اید، معتبرتر از چیزی است که ابزار حدس زده.
3. **نوشتن فقط روی سلول‌های هدف.** به‌جای بازنویسی کل ردیف، برای هر ردیف
   محدوده‌های پیوسته‌ی همان ستون‌ها ساخته و در یک درخواست فرستاده می‌شود؛
   پس ستون‌های دیگرِ همان ردیف حتی «بازنویسی با مقدار قبلی» هم نمی‌شوند.

بدون ``gspread`` هم کار می‌کند: آن‌وقت ورودی یک فایل ``csv``/``xlsx`` است و
خروجی یک فایل تازه — همان ستون‌ها، بدون نیاز به حساب گوگل.
"""

from __future__ import annotations

import csv
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Iterable, Sequence

from . import fields
from .details import label_key

#: کلیدی که مقدار ستون عنوان همیشه زیرش هم گذاشته می‌شود
TITLE_KEY = "title"

#: ستون‌هایی که فاز ۵ حق نوشتن در آن‌ها را ندارد (با نام داخلی یا عنوان ستون)
DEFAULT_NEVER_WRITE: tuple[str, ...] = ("title", "status", "product_id", "وضعیت", "شناسه محصول")


class SheetError(RuntimeError):
    """خطای قابل‌نمایش به کاربر (نه traceback)."""


# ---------------------------------------------------------------------------
# ستون‌ها
# ---------------------------------------------------------------------------


def column_letter(index: int) -> str:
    """۰ → ``A``، ۲۶ → ``AA``."""
    letters = ""
    index += 1
    while index > 0:
        index, remainder = divmod(index - 1, 26)
        letters = chr(ord("A") + remainder) + letters
    return letters


def column_index(letter: str) -> int:
    """``B`` → ۱. برای وقتی کاربر به‌جای عنوان ستون، حرف ستون را داده."""
    value = 0
    for char in str(letter).strip().upper():
        if not "A" <= char <= "Z":
            raise SheetError(f"حرف ستون نامعتبر است: {letter!r}")
        value = value * 26 + (ord(char) - ord("A") + 1)
    if value <= 0:
        raise SheetError(f"حرف ستون نامعتبر است: {letter!r}")
    return value - 1


# ---------------------------------------------------------------------------
# ردیف‌ها و به‌روزرسانی‌ها
# ---------------------------------------------------------------------------


@dataclass
class SheetRow:
    """یک ردیف شیت، فقط با ستون‌هایی که می‌شناسیم."""

    number: int  # شماره‌ی ردیف در شیت (۱-پایه)
    values: dict[str, str] = field(default_factory=dict)

    @property
    def title(self) -> str:
        return (self.values.get(TITLE_KEY) or "").strip()

    def empty_fields(self, names: Iterable[str]) -> list[str]:
        return [name for name in names if not (self.values.get(name) or "").strip()]


@dataclass
class CellUpdate:
    row_number: int
    column: int
    value: str


def group_updates(updates: Sequence[CellUpdate]) -> list[tuple[str, list[str]]]:
    """سلول‌های پراکنده → کمترین تعداد محدوده‌ی A1 پیوسته.

    ``B7=x, C7=y, K7=z`` می‌شود دو محدوده (``B7:C7`` و ``K7``) نه سه، و
    ستون‌های بین آن‌ها اصلاً در درخواست نمی‌آیند.
    """
    by_row: dict[int, dict[int, str]] = {}
    for update in updates:
        by_row.setdefault(update.row_number, {})[update.column] = update.value

    ranges: list[tuple[str, list[str]]] = []
    for row_number in sorted(by_row):
        columns = sorted(by_row[row_number])
        start = previous = columns[0]
        run = [by_row[row_number][start]]
        for column in columns[1:]:
            if column == previous + 1:
                run.append(by_row[row_number][column])
                previous = column
                continue
            ranges.append((_a1(start, previous, row_number), run))
            start = previous = column
            run = [by_row[row_number][column]]
        ranges.append((_a1(start, previous, row_number), run))
    return ranges


def _a1(first: int, last: int, row_number: int) -> str:
    if first == last:
        return f"{column_letter(first)}{row_number}"
    return f"{column_letter(first)}{row_number}:{column_letter(last)}{row_number}"


# ---------------------------------------------------------------------------
# سند گوگل‌شیت
# ---------------------------------------------------------------------------


@dataclass
class SheetSettings:
    """بخش اتصال به شیت در ``details``."""

    sheet_id: str = ""
    service_account_json: str = ""
    tab: str = ""
    lists_tab: str = "لیست‌ها"
    report_tab: str = "گزارش تکمیل"
    header_row: int = 1
    file: str = ""
    #: عنوان ستون عنوان؛ خالی = ستون اول شیت
    title_column: str = ""
    #: تنظیم دستی ستون‌ها: ``{عنوان ستون: {kind، labels، options، ...}}``
    columns: dict[str, Any] = field(default_factory=dict)
    never_write: tuple[str, ...] = DEFAULT_NEVER_WRITE
    batch_ranges: int = 400
    value_input_option: str = "RAW"

    @classmethod
    def from_mapping(cls, data: dict | None) -> "SheetSettings":
        data = data or {}
        never = data.get("never_write")
        return cls(
            sheet_id=str(data.get("sheet_id", "") or ""),
            service_account_json=str(data.get("service_account_json", "") or ""),
            tab=str(data.get("tab", "") or ""),
            lists_tab=str(data.get("lists_tab", "لیست‌ها") or ""),
            report_tab=str(data.get("report_tab", "گزارش تکمیل") or ""),
            header_row=max(1, int(data.get("header_row", 1) or 1)),
            file=str(data.get("file", "") or ""),
            title_column=str(data.get("title_column", "") or ""),
            columns=dict(data.get("columns") or {}),
            never_write=tuple(never) if never else DEFAULT_NEVER_WRITE,
            batch_ranges=max(1, int(data.get("batch_ranges", 400) or 400)),
            value_input_option=str(data.get("value_input_option", "RAW") or "RAW"),
        )


class Document:
    """رابط مشترک گوگل‌شیت و فایل محلی."""

    name = ""

    def read_grid(self) -> list[list[str]]:  # pragma: no cover - انتزاعی
        """کل شیت به‌صورت خام. تفسیر ستون‌ها کار :mod:`core.fields` است."""
        raise NotImplementedError

    def read_tab(self, title: str) -> list[list[str]]:  # pragma: no cover - انتزاعی
        return []

    def apply(self, updates: Sequence[CellUpdate]) -> int:  # pragma: no cover - انتزاعی
        raise NotImplementedError

    def write_report(self, title: str, header: Sequence[str], rows: Sequence[Sequence[Any]]) -> str:
        return ""

    def close(self) -> str:
        return ""


class GoogleSheetDocument(Document):
    """گوگل‌شیت واقعی، از راه ``gspread`` و یک service account."""

    name = "google-sheet"

    def __init__(self, settings: SheetSettings) -> None:
        self.settings = settings
        self._worksheet: Any = None
        self._spreadsheet: Any = None
        self.written_cells = 0

    # -- اتصال --------------------------------------------------------------
    @staticmethod
    def available(settings: SheetSettings) -> bool:
        if not (settings.sheet_id and settings.service_account_json):
            return False
        try:
            import importlib

            importlib.import_module("gspread")
        except ImportError:
            return False
        return Path(settings.service_account_json).exists()

    def _open(self) -> Any:  # pragma: no cover - نیازمند شبکه
        if self._worksheet is not None:
            return self._worksheet
        try:
            import gspread  # type: ignore
        except ImportError as exc:
            raise SheetError(
                "برای کار با گوگل‌شیت باید gspread نصب باشد: pip install gspread"
            ) from exc
        if not Path(self.settings.service_account_json).exists():
            raise SheetError(
                f"فایل service account پیدا نشد: {self.settings.service_account_json}"
            )
        client = gspread.service_account(filename=self.settings.service_account_json)
        try:
            self._spreadsheet = client.open_by_key(self.settings.sheet_id)
        except Exception as exc:
            raise SheetError(
                "گوگل‌شیت باز نشد. شناسه‌ی شیت را چک کنید و مطمئن شوید ایمیل "
                "service account به‌عنوان ویرایشگر به شیت دسترسی دارد.\n"
                f"پیام گوگل: {exc}"
            ) from exc
        if self.settings.tab:
            try:
                self._worksheet = self._spreadsheet.worksheet(self.settings.tab)
            except Exception as exc:
                names = [ws.title for ws in self._spreadsheet.worksheets()]
                raise SheetError(
                    f"تبِ «{self.settings.tab}» در شیت نیست. تب‌های موجود: {'، '.join(names)}"
                ) from exc
        else:
            self._worksheet = self._spreadsheet.sheet1
        return self._worksheet

    # -- خواندن -------------------------------------------------------------
    def read_grid(self) -> list[list[str]]:  # pragma: no cover - نیازمند شبکه
        worksheet = self._open()
        return [list(row) for row in worksheet.get_all_values()]

    def read_tab(self, title: str) -> list[list[str]]:  # pragma: no cover - نیازمند شبکه
        if not title:
            return []
        self._open()
        try:
            worksheet = self._spreadsheet.worksheet(title)
        except Exception:
            return []
        return [list(row) for row in worksheet.get_all_values()]

    # -- نوشتن --------------------------------------------------------------
    def apply(self, updates: Sequence[CellUpdate]) -> int:  # pragma: no cover - نیازمند شبکه
        if not updates:
            return 0
        worksheet = self._open()
        ranges = group_updates(updates)
        written = 0
        size = self.settings.batch_ranges
        for start in range(0, len(ranges), size):
            chunk = ranges[start : start + size]
            worksheet.batch_update(
                [{"range": name, "values": [values]} for name, values in chunk],
                value_input_option=self.settings.value_input_option,
            )
            written += sum(len(values) for _, values in chunk)
        self.written_cells += written
        return written

    def write_report(
        self, title: str, header: Sequence[str], rows: Sequence[Sequence[Any]]
    ) -> str:  # pragma: no cover - نیازمند شبکه
        if not title:
            return ""
        self._open()
        values = [list(header)] + [[_cell(v) for v in row] for row in rows]
        try:
            worksheet = self._spreadsheet.worksheet(title)
            worksheet.clear()
        except Exception:
            worksheet = self._spreadsheet.add_worksheet(
                title=title, rows=max(len(values) + 10, 100), cols=max(len(header) + 2, 8)
            )
        worksheet.update(values, "A1", value_input_option="RAW")
        return f"https://docs.google.com/spreadsheets/d/{self.settings.sheet_id}"

    def close(self) -> str:
        if not self.settings.sheet_id:
            return ""
        return f"https://docs.google.com/spreadsheets/d/{self.settings.sheet_id}"


class FileDocument(Document):
    """همان کار، روی یک فایل ``csv``/``xlsx`` محلی — بدون حساب گوگل.

    برای وقتی که هنوز service account نساخته‌اید یا می‌خواهید روی یک نمونه‌ی
    کوچک تست کنید: ستون‌ها را از شیت اکسپورت می‌گیرید، اینجا پر می‌شود، و
    خروجی را دوباره در شیت پیست می‌کنید.
    """

    name = "file"

    def __init__(self, path: str | Path, settings: SheetSettings) -> None:
        self.path = Path(path)
        self.settings = settings
        self.values: list[list[str]] = []
        self.output_path = self.path.with_name(f"{self.path.stem}-filled.xlsx")
        self.written_cells = 0

    def read_grid(self) -> list[list[str]]:
        if not self.path.exists():
            raise SheetError(f"فایل ورودی پیدا نشد: {self.path}")
        if self.path.suffix.lower() in {".xlsx", ".xlsm"}:
            self.values = _read_xlsx(self.path)
        else:
            self.values = _read_csv(self.path)
        return self.values

    def apply(self, updates: Sequence[CellUpdate]) -> int:
        for update in updates:
            while len(self.values) < update.row_number:
                self.values.append([])
            row = self.values[update.row_number - 1]
            while len(row) <= update.column:
                row.append("")
            row[update.column] = update.value
        self.written_cells += len(updates)
        return len(updates)

    def write_report(
        self, title: str, header: Sequence[str], rows: Sequence[Sequence[Any]]
    ) -> str:
        target = self.path.with_name(f"{self.path.stem}-report.csv")
        with target.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.writer(handle)
            writer.writerow(list(header))
            for row in rows:
                writer.writerow([_cell(value) for value in row])
        return str(target)

    def close(self) -> str:
        if not self.values:
            return ""
        try:
            from openpyxl import Workbook  # type: ignore
        except ImportError:
            target = self.path.with_name(f"{self.path.stem}-filled.csv")
            with target.open("w", encoding="utf-8-sig", newline="") as handle:
                writer = csv.writer(handle)
                for row in self.values:
                    writer.writerow(row)
            self.output_path = target
            return str(target)
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = self.settings.tab[:31] or "محصولات"
        worksheet.sheet_view.rightToLeft = True
        for row in self.values:
            worksheet.append([_cell(value) for value in row])
        workbook.save(self.output_path)
        return str(self.output_path)


def _read_csv(path: Path) -> list[list[str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return [list(row) for row in csv.reader(handle)]


def _read_xlsx(path: Path) -> list[list[str]]:
    try:
        from openpyxl import load_workbook  # type: ignore
    except ImportError as exc:  # pragma: no cover - وابسته به محیط
        raise SheetError("برای خواندن xlsx باید openpyxl نصب باشد: pip install openpyxl") from exc
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook.active
    rows = [
        ["" if cell is None else str(cell) for cell in row]
        for row in worksheet.iter_rows(values_only=True)
    ]
    workbook.close()
    return rows


def plan_from_grid(
    grid: Sequence[Sequence[str]],
    settings: SheetSettings,
    lists: dict[str, list[str]] | None = None,
) -> fields.FieldPlan:
    """سطر عنوانِ شیت → نقشه‌ی ستون‌ها.

    این تنها جایی است که تصمیم گرفته می‌شود «این شیت چه ستون‌هایی دارد»، و
    تصمیمش از خودِ شیت می‌آید نه از کد: عنوان هر ستون، هم برچسب جستجو در
    منابع می‌شود و هم نوع فیلد را تعیین می‌کند.
    """
    if not grid:
        raise SheetError("شیت خالی است.")
    header_index = min(settings.header_row, len(grid)) - 1
    plan = fields.build_plan(
        grid[header_index],
        lists=lists or {},
        overrides=settings.columns,
        title_column=settings.title_column,
    )
    if not plan.specs:
        raise SheetError(
            f"سطر {settings.header_row} شیت عنوان ستون ندارد."
            " اگر عنوان‌ها در سطر دیگری‌اند، details.header_row را عوض کنید."
        )
    return plan


def rows_from_grid(
    grid: Sequence[Sequence[str]], plan: fields.FieldPlan, header_row: int = 1
) -> list[SheetRow]:
    """ردیف‌های شیت با مقدار هر ستونِ نقشه‌شده."""
    header_index = min(header_row, len(grid)) - 1
    title_key = plan.title_key
    rows: list[SheetRow] = []
    for offset, raw in enumerate(grid[header_index + 1 :], start=header_index + 2):
        cells = list(raw)
        values = {
            key: (cells[index].strip() if index < len(cells) and cells[index] else "")
            for key, index in plan.columns.items()
        }
        if title_key:
            values[TITLE_KEY] = values.get(title_key, "")
        if not (values.get(TITLE_KEY) or "").strip():
            continue  # ردیف خالی وسط شیت، پایان کار نیست
        rows.append(SheetRow(number=offset, values=values))
    return rows


def _cell(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, (list, tuple, set)):
        return "، ".join(str(item) for item in value)
    return value


def open_document(settings: SheetSettings) -> Document:
    """سند مناسب را باز می‌کند: گوگل‌شیت اگر تنظیم شده، وگرنه فایل محلی."""
    if settings.sheet_id:
        if not GoogleSheetDocument.available(settings):
            raise SheetError(
                "گوگل‌شیت تنظیم شده ولی قابل استفاده نیست. سه چیز لازم است:\n"
                "  ۱. نصب کتابخانه:  pip install gspread\n"
                "  ۲. مسیر درست فایل service account در details.service_account_json\n"
                "  ۳. اشتراک شیت با ایمیل همان service account (دسترسی ویرایش)"
            )
        return GoogleSheetDocument(settings)
    if settings.file:
        return FileDocument(settings.file, settings)
    raise SheetError(
        "نه شناسه‌ی گوگل‌شیت داده شده و نه فایل ورودی."
        " در تب «تکمیل دیتیل» پنل یکی را پر کنید یا details.sheet_id / details.file را در config بگذارید."
    )
